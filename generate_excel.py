import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_test_excel():
    wb = openpyxl.Workbook()
    
    # Sheet 1: 导入数据 (车满满批量下单标准格式)
    ws = wb.active
    ws.title = "导入数据"
    
    headers = [
        "schema_version",
        "batch_id",
        "source_record_id",
        "source_label",
        "destination_text",
        "delivery_type",
        "sender_name",
        "receiver_name",
        "receiver_mobile",
        "goods_name",
        "package",
        "quantity",
        "weight",
        "volume",
        "freight",
        "payment_type"
    ]
    
    data = [
        # --- 第一组：发货商、收货人、电话相同 (用例 1 & 2) ---
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_001",
            "source_label": "用例01_同组A-1",
            "destination_text": "通海县",
            "delivery_type": "delivery",
            "sender_name": "上海恒通实业发展有限公司",
            "receiver_name": "张伟",
            "receiver_mobile": "13800138001",
            "goods_name": "机械配件",
            "package": "纸箱",
            "quantity": 10,
            "weight": 45.5,
            "volume": 0.35,
            "freight": 180.00,
            "payment_type": "pay_billing"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_002",
            "source_label": "用例02_同组A-2",
            "destination_text": "玉溪市",
            "delivery_type": "pickup",
            "sender_name": "上海恒通实业发展有限公司",
            "receiver_name": "张伟",
            "receiver_mobile": "13800138001",
            "goods_name": "工业五金",
            "package": "木架",
            "quantity": 15,
            "weight": 68.0,
            "volume": 0.52,
            "freight": 260.00,
            "payment_type": "pay_arrival"
        },
        # --- 第二组：发货商、收货人、电话相同 (用例 3 & 4) ---
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_003",
            "source_label": "用例03_同组B-1",
            "destination_text": "通海县",
            "delivery_type": "delivery",
            "sender_name": "广州盛达供应链有限公司",
            "receiver_name": "李秀英",
            "receiver_mobile": "13911223344",
            "goods_name": "电子元件",
            "package": "纸箱",
            "quantity": 8,
            "weight": 22.0,
            "volume": 0.18,
            "freight": 120.50,
            "payment_type": "pay_billing"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_004",
            "source_label": "用例04_同组B-2",
            "destination_text": "玉溪市",
            "delivery_type": "pickup",
            "sender_name": "广州盛达供应链有限公司",
            "receiver_name": "李秀英",
            "receiver_mobile": "13911223344",
            "goods_name": "仪器仪表",
            "package": "木箱",
            "quantity": 4,
            "weight": 35.0,
            "volume": 0.28,
            "freight": 195.00,
            "payment_type": "pay_receipt"
        },
        # --- 其余 8 组：全部互不相同 (用例 5 ~ 12) ---
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_005",
            "source_label": "用例05_独立01",
            "destination_text": "西安市",
            "delivery_type": "delivery",
            "sender_name": "北京中科智控科技有限公司",
            "receiver_name": "王磊",
            "receiver_mobile": "13712345678",
            "goods_name": "传感器模块",
            "package": "托盘",
            "quantity": 6,
            "weight": 18.5,
            "volume": 0.15,
            "freight": 150.00,
            "payment_type": "pay_billing"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_006",
            "source_label": "用例06_独立02",
            "destination_text": "南京市",
            "delivery_type": "pickup",
            "sender_name": "深圳市捷达兴物流设备厂",
            "receiver_name": "赵雪梅",
            "receiver_mobile": "13698765432",
            "goods_name": "输送带滚筒",
            "package": "编织袋",
            "quantity": 20,
            "weight": 110.0,
            "volume": 0.85,
            "freight": 380.00,
            "payment_type": "pay_arrival"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_007",
            "source_label": "用例07_独立03",
            "destination_text": "郑州市",
            "delivery_type": "delivery",
            "sender_name": "浙江吉美日用品制造厂",
            "receiver_name": "陈建国",
            "receiver_mobile": "13588990011",
            "goods_name": "家居收纳盒",
            "package": "纸箱",
            "quantity": 30,
            "weight": 55.0,
            "volume": 1.20,
            "freight": 290.00,
            "payment_type": "pay_receipt"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_008",
            "source_label": "用例08_独立04",
            "destination_text": "长沙市",
            "delivery_type": "pickup",
            "sender_name": "江苏苏博特新材料有限公司",
            "receiver_name": "刘芳",
            "receiver_mobile": "15866778899",
            "goods_name": "涂料辅料",
            "package": "铁桶",
            "quantity": 12,
            "weight": 240.0,
            "volume": 0.60,
            "freight": 450.00,
            "payment_type": "pay_billing"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_009",
            "source_label": "用例09_独立05",
            "destination_text": "江津区",
            "delivery_type": "delivery",
            "sender_name": "山东鲁泰纺织服装有限公司",
            "receiver_name": "孙宏伟",
            "receiver_mobile": "15900223344",
            "goods_name": "纯棉工作服",
            "package": "胶框",
            "quantity": 25,
            "weight": 75.0,
            "volume": 0.90,
            "freight": 310.00,
            "payment_type": "pay_arrival"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_010",
            "source_label": "用例10_独立06",
            "destination_text": "顺义区",
            "delivery_type": "pickup",
            "sender_name": "福建万达汽车配件有限公司",
            "receiver_name": "周美玲",
            "receiver_mobile": "18611229988",
            "goods_name": "刹车片",
            "package": "缠绕膜",
            "quantity": 18,
            "weight": 85.0,
            "volume": 0.45,
            "freight": 230.00,
            "payment_type": "pay_receipt"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_011",
            "source_label": "用例11_独立07",
            "destination_text": "昆山市",
            "delivery_type": "delivery",
            "sender_name": "四川德胜生物制品有限公司",
            "receiver_name": "吴德华",
            "receiver_mobile": "18955667788",
            "goods_name": "试验试剂",
            "package": "木箱",
            "quantity": 5,
            "weight": 15.0,
            "volume": 0.12,
            "freight": 210.00,
            "payment_type": "pay_billing"
        },
        {
            "schema_version": "v1.1",
            "batch_id": "BATCH_20260815_01",
            "source_record_id": "TC_012",
            "source_label": "用例12_独立08",
            "destination_text": "济南市",
            "delivery_type": "delivery",
            "sender_name": "河北冀南电线电缆有限公司",
            "receiver_name": "黄丽",
            "receiver_mobile": "17733445566",
            "goods_name": "绝缘铜线",
            "package": "托盘",
            "quantity": 16,
            "weight": 160.0,
            "volume": 0.70,
            "freight": 420.00,
            "payment_type": "pay_arrival"
        }
    ]
    
    # 写入表头
    ws.append(headers)
    
    # 写入数据
    for item in data:
        row = [item[h] for h in headers]
        ws.append(row)
        
    # 美化样式
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    
    group_a_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # 浅绿
    group_b_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # 浅橙
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=len(data)+1, min_col=1, max_col=len(headers)), 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx, cell in enumerate(col, 1):
            cell.border = thin_border
            cell.font = Font(name="Microsoft YaHei", size=10)
            
            if row_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                # 高亮两组相同发货/收货/电话的行
                if row_idx in (2, 3):  # TC_001, TC_002
                    cell.fill = group_a_fill
                elif row_idx in (4, 5):  # TC_003, TC_004
                    cell.fill = group_b_fill
                
                # 对齐方式
                if headers[col_idx-1] in ["quantity", "weight", "volume", "freight"]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif headers[col_idx-1] in ["schema_version", "batch_id", "source_record_id", "delivery_type", "receiver_mobile", "payment_type"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    
            val_str = str(cell.value or "")
            # 计算包含中文的字符串显示宽度
            char_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if char_len > max_len:
                max_len = char_len
                
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    ws.row_dimensions[1].height = 26
    for r in range(2, len(data) + 2):
        ws.row_dimensions[r].height = 22
        
    # 保存文件
    output_path = r"d:\Goings\APPProjects\Chemanman\测试开单数据_12条.xlsx"
    wb.save(output_path)
    print(f"Excel generated successfully at: {output_path}")

if __name__ == "__main__":
    generate_test_excel()
